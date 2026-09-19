"""
Reports Tab — التقارير والإحصائيات.
+ Light/Dark theme support (reads from Theme).

- فلتر بالتاريخ (من/إلى) — آخر 30 يوم افتراضيًا
- 6 كروت إحصائية (إيرادات، حجوزات، عملاء، كاش، كارت، محفظة)
- 4 تابات داخلية:
    🏆 أعلى الخدمات
    👨‍💼 إيرادات الحلاقين
    📅 أكثر الأيام ازدحامًا
    😴 عملاء غير نشطين
- تصدير Excel
"""

import os
from datetime import datetime, timedelta

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QGridLayout, QLabel,
    QPushButton, QFrame, QTableWidget, QTableWidgetItem,
    QHeaderView, QMessageBox, QAbstractItemView, QTabWidget,
    QLineEdit, QSizePolicy
)
from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QFont

from theme import Theme


try:
    import openpyxl
    from openpyxl.styles import Font as XLFont
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ============================================================
# Stat Card
# ============================================================
class StatCard(QFrame):
    """بطاقة إحصائية صغيرة — theme-aware."""

    def __init__(self, title, value="—", icon="📊",
                 color_key="primary", parent=None):
        super().__init__(parent)

        color = Theme.color(color_key)

        self.setStyleSheet(
            f"QFrame {{ background-color: {Theme.color('surface')}; "
            f"border: 1px solid {color}; border-radius: 12px; }}"
        )
        self.setMinimumHeight(100)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 12, 16, 12)
        layout.setSpacing(4)

        title_row = QHBoxLayout()
        title_row.setSpacing(6)

        icon_lbl = QLabel(icon)
        icon_lbl.setStyleSheet(
            f"font-size: 18px; color: {color}; background: transparent;"
        )
        title_row.addWidget(icon_lbl)

        title_lbl = QLabel(title)
        title_lbl.setStyleSheet(
            f"color: {Theme.color('text_muted')}; font-size: 11px; "
            f"font-weight: bold; background: transparent;"
        )
        title_row.addWidget(title_lbl)
        title_row.addStretch()
        layout.addLayout(title_row)

        self.value_label = QLabel(str(value))
        self.value_label.setStyleSheet(
            f"color: {color}; font-size: 22px; font-weight: bold; "
            f"background: transparent; padding: 2px;"
        )
        self.value_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        layout.addWidget(self.value_label)

    def set_value(self, value):
        self.value_label.setText(str(value))


# ============================================================
# Reports Tab
# ============================================================
class ReportsTab(QWidget):
    """تاب التقارير والإحصائيات."""

    def __init__(self, db, user):
        super().__init__()
        self.db = db
        self.user = user

        self._build_ui()
        self.refresh()

    # ============================================
    # UI
    # ============================================
    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(14)

        # Header
        header = QHBoxLayout()
        title = QLabel("📊  التقارير والإحصائيات")
        title.setStyleSheet(
            f"font-size: 22px; font-weight: bold; "
            f"color: {Theme.color('primary')}; padding: 4px; "
            f"background: transparent;"
        )
        header.addWidget(title)
        header.addStretch()
        layout.addLayout(header)

        # Filters Card
        filters_card = QFrame()
        filters_card.setStyleSheet(
            f"QFrame {{ background-color: {Theme.color('surface')}; "
            f"border: 1px solid {Theme.color('border')}; border-radius: 12px; }}"
        )
        filters_layout = QHBoxLayout(filters_card)
        filters_layout.setContentsMargins(18, 14, 18, 14)
        filters_layout.setSpacing(10)

        filters_layout.addWidget(self._lbl("من:"))

        default_from = (datetime.now() - timedelta(days=30)).strftime("%Y-%m-%d")
        default_to = datetime.now().strftime("%Y-%m-%d")

        self.date_from = QLineEdit(default_from)
        self.date_from.setPlaceholderText("YYYY-MM-DD")
        self.date_from.setMinimumHeight(40)
        self.date_from.setMaximumWidth(150)
        self.date_from.setAlignment(Qt.AlignCenter)
        filters_layout.addWidget(self.date_from)

        filters_layout.addWidget(self._lbl("إلى:"))

        self.date_to = QLineEdit(default_to)
        self.date_to.setPlaceholderText("YYYY-MM-DD")
        self.date_to.setMinimumHeight(40)
        self.date_to.setMaximumWidth(150)
        self.date_to.setAlignment(Qt.AlignCenter)
        filters_layout.addWidget(self.date_to)

        # Quick range buttons
        for label, days in [("اليوم", 0), ("أسبوع", 7), ("شهر", 30),
                             ("3 شهور", 90), ("سنة", 365)]:
            btn = QPushButton(label)
            btn.setMinimumHeight(40)
            btn.setCursor(Qt.PointingHandCursor)
            btn.setStyleSheet(
                f"QPushButton {{ background-color: {Theme.color('surface_alt')}; "
                f"color: {Theme.color('text')}; border: none; "
                f"border-radius: 6px; padding: 6px 14px; "
                f"font-size: 11px; font-weight: bold; }}"
                f"QPushButton:hover {{ background-color: {Theme.color('border_strong')}; }}"
            )
            btn.clicked.connect(lambda _, d=days: self._quick_range(d))
            filters_layout.addWidget(btn)

        filters_layout.addStretch()

        apply_btn = QPushButton("🔍  تطبيق")
        apply_btn.setMinimumHeight(40)
        apply_btn.setMinimumWidth(120)
        apply_btn.setCursor(Qt.PointingHandCursor)
        apply_btn.setStyleSheet(
            f"QPushButton {{ background-color: {Theme.color('primary')}; "
            f"color: white; border: none; border-radius: 8px; "
            f"font-weight: bold; font-size: 12px; padding: 6px 14px; }}"
            f"QPushButton:hover {{ background-color: {Theme.color('primary_hover')}; }}"
        )
        apply_btn.clicked.connect(self.refresh)
        filters_layout.addWidget(apply_btn)

        export_btn = QPushButton("📊  تصدير Excel")
        export_btn.setMinimumHeight(40)
        export_btn.setMinimumWidth(150)
        export_btn.setCursor(Qt.PointingHandCursor)
        export_btn.setStyleSheet(
            f"QPushButton {{ background-color: {Theme.color('success')}; "
            f"color: white; border: none; border-radius: 8px; "
            f"font-weight: bold; font-size: 12px; padding: 6px 14px; }}"
            f"QPushButton:hover {{ background-color: {Theme.color('success_hover')}; }}"
        )
        export_btn.clicked.connect(self._export_excel)
        filters_layout.addWidget(export_btn)

        layout.addWidget(filters_card)

        # Stat Cards (6)
        cards_grid = QGridLayout()
        cards_grid.setSpacing(10)

        self.card_revenue = StatCard("إجمالي الإيرادات", "0 ج", "💰", "success")
        self.card_bookings = StatCard("عدد الحجوزات", "0", "📅", "primary")
        self.card_customers = StatCard("عدد العملاء", "0", "👥", "info")
        self.card_cash = StatCard("كاش", "0 ج", "💵", "success")
        self.card_card = StatCard("فيزا / كارت", "0 ج", "💳", "primary")
        self.card_wallet = StatCard("محفظة", "0 ج", "📱", "warning")

        cards_grid.addWidget(self.card_revenue, 0, 0)
        cards_grid.addWidget(self.card_bookings, 0, 1)
        cards_grid.addWidget(self.card_customers, 0, 2)
        cards_grid.addWidget(self.card_cash, 1, 0)
        cards_grid.addWidget(self.card_card, 1, 1)
        cards_grid.addWidget(self.card_wallet, 1, 2)

        layout.addLayout(cards_grid)

        # Sub-tabs
        self.sub_tabs = QTabWidget()
        self.sub_tabs.setStyleSheet(f"""
            QTabWidget::pane {{
                border: 1px solid {Theme.color('border')};
                border-radius: 8px;
                background-color: {Theme.color('bg')};
                padding: 12px;
            }}
            QTabBar::tab {{
                background-color: {Theme.color('surface')};
                color: {Theme.color('text_muted')};
                padding: 10px 22px;
                margin-right: 4px;
                border-top-left-radius: 8px;
                border-top-right-radius: 8px;
                font-size: 12px;
                font-weight: bold;
                min-width: 130px;
            }}
            QTabBar::tab:selected {{
                background-color: {Theme.color('primary')};
                color: #ffffff;
            }}
            QTabBar::tab:hover:!selected {{
                background-color: {Theme.color('surface_alt')};
                color: {Theme.color('text')};
            }}
        """)

        self.tab_services = self._build_table_tab(
            ["الخدمة", "عدد المرات", "الإجمالي"]
        )
        self.sub_tabs.addTab(self.tab_services, "  🏆  أعلى الخدمات  ")

        self.tab_barbers = self._build_table_tab(
            ["الحلاق", "عدد الحجوزات", "الإجمالي"]
        )
        self.sub_tabs.addTab(self.tab_barbers, "  👨‍💼  إيرادات الحلاقين  ")

        self.tab_weekdays = self._build_table_tab(
            ["اليوم", "عدد الحجوزات"]
        )
        self.sub_tabs.addTab(self.tab_weekdays, "  📅  أكثر الأيام  ")

        self.tab_inactive = self._build_table_tab(
            ["العميل", "التليفون", "آخر زيارة"]
        )
        self.sub_tabs.addTab(self.tab_inactive, "  😴  عملاء غير نشطين  ")

        layout.addWidget(self.sub_tabs, stretch=1)

    def _lbl(self, text):
        lbl = QLabel(text)
        lbl.setStyleSheet(
            f"color: {Theme.color('text')}; font-size: 12px; "
            f"font-weight: bold; background: transparent; padding: 0 4px;"
        )
        return lbl

    def _build_table_tab(self, headers):
        widget = QWidget()
        w_layout = QVBoxLayout(widget)
        w_layout.setContentsMargins(0, 0, 0, 0)

        table = QTableWidget()
        table.setColumnCount(len(headers))
        table.setHorizontalHeaderLabels(headers)
        table.verticalHeader().setVisible(False)
        table.setAlternatingRowColors(True)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setSelectionMode(QAbstractItemView.SingleSelection)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setShowGrid(False)
        table.setMinimumHeight(240)
        table.setLayoutDirection(Qt.RightToLeft)

        h = table.horizontalHeader()
        for i in range(len(headers)):
            h.setSectionResizeMode(i, QHeaderView.Stretch)

        w_layout.addWidget(table)
        widget._table = table
        return widget

    # ============================================
    # Quick range
    # ============================================
    def _quick_range(self, days):
        today = datetime.now()
        if days == 0:
            self.date_from.setText(today.strftime("%Y-%m-%d"))
            self.date_to.setText(today.strftime("%Y-%m-%d"))
        else:
            self.date_from.setText(
                (today - timedelta(days=days)).strftime("%Y-%m-%d")
            )
            self.date_to.setText(today.strftime("%Y-%m-%d"))
        self.refresh()

    # ============================================
    # Refresh
    # ============================================
    def refresh(self):
        date_from = self.date_from.text().strip() or None
        date_to = self.date_to.text().strip() or None

        try:
            if date_from:
                datetime.strptime(date_from, "%Y-%m-%d")
            if date_to:
                datetime.strptime(date_to, "%Y-%m-%d")
        except ValueError:
            QMessageBox.warning(
                self, "خطأ",
                "صيغة التاريخ غلط. استخدم YYYY-MM-DD"
            )
            return

        try:
            rev = self.db.revenue_report_range(date_from, date_to)
            total_rev = rev["total_revenue"]
            total_bookings = rev["total_bookings"]

            total_customers = self.db.count_customers_range(date_from, date_to)

            pm = self.db.payment_breakdown_range(date_from, date_to)
            cash = pm["cash"]["total"]
            card = pm["card"]["total"]
            wallet = pm["wallet"]["total"]

            services = self.db.popular_services_report_range(date_from, date_to)
            barbers = self.db.revenue_by_barber_range(date_from, date_to)
            weekdays = self.db.busiest_weekdays_range(date_from, date_to)
            inactive = self.db.inactive_customers_range(date_from, date_to)

        except Exception as e:
            QMessageBox.critical(self, "خطأ", f"فشل تحميل التقارير:\n{e}")
            return

        self.card_revenue.set_value(f"{total_rev:,.2f} ج")
        self.card_bookings.set_value(f"{total_bookings:,}")
        self.card_customers.set_value(f"{total_customers:,}")
        self.card_cash.set_value(f"{cash:,.2f} ج")
        self.card_card.set_value(f"{card:,.2f} ج")
        self.card_wallet.set_value(f"{wallet:,.2f} ج")

        self._fill_services_table(services)
        self._fill_barbers_table(barbers)
        self._fill_weekdays_table(weekdays)
        self._fill_inactive_table(inactive)

    def _fill_services_table(self, services):
        table = self.tab_services._table
        table.setRowCount(len(services))
        green = QColor(Theme.color("success"))
        text_color = QColor(Theme.color("text"))

        for row, (name, count, total) in enumerate(services):
            i0 = QTableWidgetItem(str(name))
            i0.setTextAlignment(Qt.AlignCenter)
            i0.setForeground(text_color)
            table.setItem(row, 0, i0)

            i1 = QTableWidgetItem(str(count))
            i1.setTextAlignment(Qt.AlignCenter)
            i1.setForeground(text_color)
            table.setItem(row, 1, i1)

            i2 = QTableWidgetItem(f"{total or 0:,.2f} ج")
            i2.setTextAlignment(Qt.AlignCenter)
            i2.setForeground(green)
            table.setItem(row, 2, i2)

    def _fill_barbers_table(self, barbers):
        table = self.tab_barbers._table
        table.setRowCount(len(barbers))
        green = QColor(Theme.color("success"))
        blue = QColor(Theme.color("primary"))
        text_color = QColor(Theme.color("text"))

        for row, (name, count, total) in enumerate(barbers):
            i0 = QTableWidgetItem(str(name))
            i0.setTextAlignment(Qt.AlignCenter)
            i0.setForeground(text_color)
            table.setItem(row, 0, i0)

            i1 = QTableWidgetItem(str(count))
            i1.setTextAlignment(Qt.AlignCenter)
            i1.setForeground(blue)
            table.setItem(row, 1, i1)

            i2 = QTableWidgetItem(f"{total or 0:,.2f} ج")
            i2.setTextAlignment(Qt.AlignCenter)
            i2.setForeground(green)
            table.setItem(row, 2, i2)

    def _fill_weekdays_table(self, weekdays):
        table = self.tab_weekdays._table
        table.setRowCount(len(weekdays))

        sorted_wd = sorted(weekdays, key=lambda x: x[1], reverse=True)
        max_count = max((c for _, c in sorted_wd), default=1) or 1

        text_color = QColor(Theme.color("text"))
        red = QColor(Theme.color("danger"))
        orange = QColor(Theme.color("warning"))
        green = QColor(Theme.color("success"))

        for row, (name, count) in enumerate(sorted_wd):
            i0 = QTableWidgetItem(str(name))
            i0.setTextAlignment(Qt.AlignCenter)
            i0.setForeground(text_color)
            table.setItem(row, 0, i0)

            i1 = QTableWidgetItem(str(count))
            i1.setTextAlignment(Qt.AlignCenter)
            ratio = count / max_count if max_count else 0
            if ratio > 0.7:
                i1.setForeground(red)
            elif ratio > 0.4:
                i1.setForeground(orange)
            else:
                i1.setForeground(green)
            table.setItem(row, 1, i1)

    def _fill_inactive_table(self, inactive):
        table = self.tab_inactive._table
        table.setRowCount(len(inactive))
        gray = QColor(Theme.color("text_muted"))
        orange = QColor(Theme.color("warning"))
        text_color = QColor(Theme.color("text"))

        for row, (cid, name, phone, last_visit) in enumerate(inactive):
            i0 = QTableWidgetItem(str(name))
            i0.setTextAlignment(Qt.AlignCenter)
            i0.setForeground(text_color)
            table.setItem(row, 0, i0)

            i1 = QTableWidgetItem(str(phone or "—"))
            i1.setTextAlignment(Qt.AlignCenter)
            i1.setForeground(text_color)
            table.setItem(row, 1, i1)

            i2 = QTableWidgetItem(str(last_visit))
            i2.setTextAlignment(Qt.AlignCenter)
            if "لم يحجز" in str(last_visit):
                i2.setForeground(orange)
            else:
                i2.setForeground(gray)
            table.setItem(row, 2, i2)

    # ============================================
    # Export Excel
    # ============================================
    def _export_excel(self):
        if not HAS_OPENPYXL:
            QMessageBox.critical(
                self, "غير متاح",
                "التصدير لـ Excel محتاج مكتبة openpyxl.\n\n"
                "ثبّتها بالأمر:\n"
                "pip install openpyxl"
            )
            return

        date_from = self.date_from.text().strip() or None
        date_to = self.date_to.text().strip() or None

        try:
            wb = openpyxl.Workbook()

            ws1 = wb.active
            ws1.title = "الملخص"
            ws1.append(["تقرير التقارير"])
            ws1.append(["من", date_from or "البداية"])
            ws1.append(["إلى", date_to or "النهاية"])
            ws1.append([])
            ws1.append(["البند", "القيمة"])

            rev = self.db.revenue_report_range(date_from, date_to)
            total_customers = self.db.count_customers_range(date_from, date_to)
            pm = self.db.payment_breakdown_range(date_from, date_to)

            ws1.append(["إجمالي الإيرادات", f"{rev['total_revenue']:,.2f} ج"])
            ws1.append(["عدد الحجوزات", rev["total_bookings"]])
            ws1.append(["عدد العملاء", total_customers])
            ws1.append(["كاش", f"{pm['cash']['total']:,.2f} ج"])
            ws1.append(["فيزا / كارت", f"{pm['card']['total']:,.2f} ج"])
            ws1.append(["محفظة", f"{pm['wallet']['total']:,.2f} ج"])

            for cell in ws1["A"]:
                cell.font = XLFont(bold=True)

            ws2 = wb.create_sheet("أعلى الخدمات")
            ws2.append(["الخدمة", "عدد المرات", "الإجمالي"])
            for cell in ws2[1]:
                cell.font = XLFont(bold=True)
            for name, count, total in self.db.popular_services_report_range(
                date_from, date_to
            ):
                ws2.append([name, count, total or 0])

            ws3 = wb.create_sheet("إيرادات الحلاقين")
            ws3.append(["الحلاق", "عدد الحجوزات", "الإجمالي"])
            for cell in ws3[1]:
                cell.font = XLFont(bold=True)
            for name, count, total in self.db.revenue_by_barber_range(
                date_from, date_to
            ):
                ws3.append([name, count, total or 0])

            ws4 = wb.create_sheet("أكثر الأيام")
            ws4.append(["اليوم", "عدد الحجوزات"])
            for cell in ws4[1]:
                cell.font = XLFont(bold=True)
            for name, count in self.db.busiest_weekdays_range(
                date_from, date_to
            ):
                ws4.append([name, count])

            ws5 = wb.create_sheet("عملاء غير نشطين")
            ws5.append(["العميل", "التليفون", "آخر زيارة"])
            for cell in ws5[1]:
                cell.font = XLFont(bold=True)
            for cid, name, phone, last_visit in self.db.inactive_customers_range(
                date_from, date_to
            ):
                ws5.append([name, phone or "", last_visit])

            for ws in [ws1, ws2, ws3, ws4, ws5]:
                for col in ws.columns:
                    max_len = 0
                    col_letter = col[0].column_letter
                    for cell in col:
                        if cell.value is not None:
                            max_len = max(max_len, len(str(cell.value)))
                    ws.column_dimensions[col_letter].width = min(max_len + 3, 50)

            reports_dir = os.path.join(self.db.base_dir, "reports")
            os.makedirs(reports_dir, exist_ok=True)
            stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
            file_path = os.path.join(reports_dir, f"report_{stamp}.xlsx")
            wb.save(file_path)

            if os.name == "nt":
                os.startfile(file_path)

            QMessageBox.information(
                self, "تم",
                f"تم حفظ التقرير في:\n{file_path}"
            )
        except Exception as e:
            QMessageBox.critical(self, "خطأ", f"فشل التصدير:\n{e}")